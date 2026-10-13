
import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from io import BytesIO
import json
import os

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# =========================================================
# 1. CONFIGURACIÓN GENERAL
# =========================================================

st.set_page_config(
    page_title="Frappés POS | Sistema Ejecutivo",
    page_icon="🥤",
    layout="wide",
    initial_sidebar_state="expanded"
)

TZ = ZoneInfo("America/La_Paz")
DB = Path(__file__).resolve().parent / "frappes_database.xlsx"

st.markdown("""
<style>
.stApp {
    background-color: #0b1220;
    color: #f1f5f9;
}
[data-testid="stMetric"] {
    background: linear-gradient(135deg,#172554,#164e63);
    border: 1px solid #334155;
    padding: 16px;
    border-radius: 14px;
}
[data-testid="stMetricLabel"] {
    color: #cbd5e1;
}
[data-testid="stMetricValue"] {
    color: #67e8f9;
    font-weight: 800;
}
h1, h2, h3 {
    color: #67e8f9;
}
.stButton > button {
    border-radius: 9px;
    font-weight: 700;
}
</style>
""", unsafe_allow_html=True)


def ahora():
    return datetime.now(TZ)


# =========================================================
# 2. INVENTARIO INICIAL
# =========================================================

INVENTARIO_INICIAL = {
    "Frappé Clásico Café": {
        "categoria": "Frappés",
        "precio": 20.0,
        "costo": 7.0,
        "stock": 50
    },
    "Frappé Oreo Deluxe": {
        "categoria": "Frappés",
        "precio": 25.0,
        "costo": 9.0,
        "stock": 40
    },
    "Frappé Caramel Macchiato": {
        "categoria": "Frappés",
        "precio": 24.0,
        "costo": 8.5,
        "stock": 30
    },
    "Frappé Matcha Premium": {
        "categoria": "Frappés",
        "precio": 28.0,
        "costo": 10.0,
        "stock": 25
    },
    "Bobas de Tapioca": {
        "categoria": "Extras",
        "precio": 5.0,
        "costo": 1.5,
        "stock": 100
    },
    "Bobas Explosivas Frutilla": {
        "categoria": "Extras",
        "precio": 6.0,
        "costo": 2.0,
        "stock": 80
    },
    "Chantilly Extra": {
        "categoria": "Extras",
        "precio": 3.0,
        "costo": 1.0,
        "stock": 60
    }
}


# =========================================================
# 3. BASE DE DATOS LOCAL EN EXCEL
# =========================================================

COLUMNAS_VENTAS = [
    "ID", "Fecha", "Hora", "Cliente", "Modalidad",
    "Metodo_Pago", "Notas", "Total_Venta", "Costo_Total",
    "Ganancia_Neta", "Estado", "Detalles_JSON"
]

COLUMNAS_DETALLES = [
    "ID_Pedido", "Fecha", "Hora", "Cliente",
    "Modalidad", "Metodo_Pago", "Producto", "Categoria",
    "Cantidad", "Precio_Unitario", "Costo_Unitario",
    "Ingreso_Total", "Costo_Total", "Ganancia_Neta"
]


def cargar_db():
    if not DB.exists():
        return INVENTARIO_INICIAL.copy(), [], 1

    try:
        inventario = {}
        ventas = []

        xls = pd.ExcelFile(DB)

        if "Inventario" in xls.sheet_names:
            df_inv = pd.read_excel(DB, sheet_name="Inventario")

            for _, r in df_inv.iterrows():
                nombre = str(r["Producto"])

                if nombre and nombre != "nan":
                    inventario[nombre] = {
                        "categoria": str(r.get("Categoria", "Frappés")),
                        "precio": float(r.get("Precio", 0)),
                        "costo": float(r.get("Costo", 0)),
                        "stock": int(r.get("Stock", 0))
                    }

        if "Ventas" in xls.sheet_names:
            df_ventas = pd.read_excel(DB, sheet_name="Ventas")

            for _, r in df_ventas.iterrows():
                try:
                    detalles = json.loads(str(r.get("Detalles_JSON", "[]")))

                    ventas.append({
                        "id": int(r["ID"]),
                        "fecha": str(r["Fecha"])[:10],
                        "hora": str(r["Hora"]),
                        "cliente": str(r["Cliente"]),
                        "servicio": str(r["Modalidad"]),
                        "pago": str(r["Metodo_Pago"]),
                        "notas": str(r.get("Notas", "")),
                        "total_venta": float(r["Total_Venta"]),
                        "costo_total": float(r["Costo_Total"]),
                        "ganancia_neta": float(r["Ganancia_Neta"]),
                        "estado": str(r["Estado"]),
                        "detalles": detalles
                    })
                except (ValueError, TypeError, KeyError):
                    continue

        if not inventario:
            inventario = INVENTARIO_INICIAL.copy()

        siguiente_id = max(
            [v["id"] for v in ventas] + [0]
        ) + 1

        return inventario, ventas, siguiente_id

    except Exception as e:
        st.error(
            f"No se pudo leer la base de datos: {e}. "
            "Cierra el archivo Excel si está abierto y vuelve a intentar."
        )
        return INVENTARIO_INICIAL.copy(), [], 1


def guardar_db():
    inventario = st.session_state.inventario
    ventas = st.session_state.ventas

    df_inv = pd.DataFrame([
        {
            "Producto": nombre,
            "Categoria": p["categoria"],
            "Precio": p["precio"],
            "Costo": p["costo"],
            "Stock": p["stock"]
        }
        for nombre, p in inventario.items()
    ])

    filas_ventas = []
    filas_detalles = []

    for v in ventas:
        filas_ventas.append({
            "ID": v["id"],
            "Fecha": v["fecha"],
            "Hora": v["hora"],
            "Cliente": v["cliente"],
            "Modalidad": v["servicio"],
            "Metodo_Pago": v["pago"],
            "Notas": v.get("notas", ""),
            "Total_Venta": v["total_venta"],
            "Costo_Total": v["costo_total"],
            "Ganancia_Neta": v["ganancia_neta"],
            "Estado": v["estado"],
            "Detalles_JSON": json.dumps(
                v["detalles"], ensure_ascii=False
            )
        })

        for item in v.get("detalles", []):
            filas_detalles.append({
                "ID_Pedido": v["id"],
                "Fecha": v["fecha"],
                "Hora": v["hora"],
                "Cliente": v["cliente"],
                "Modalidad": v["servicio"],
                "Metodo_Pago": v["pago"],
                "Producto": item["producto"],
                "Categoria": item["categoria"],
                "Cantidad": item["cantidad"],
                "Precio_Unitario": item["precio"],
                "Costo_Unitario": item["costo"],
                "Ingreso_Total": item["subtotal"],
                "Costo_Total": item["costo_total"],
                "Ganancia_Neta": item["ganancia_item"]
            })

    df_ventas = pd.DataFrame(filas_ventas, columns=COLUMNAS_VENTAS)
    df_detalles = pd.DataFrame(
        filas_detalles, columns=COLUMNAS_DETALLES
    )

    temporal = DB.with_name("frappes_database_temp.xlsx")

    try:
        with pd.ExcelWriter(temporal, engine="openpyxl") as writer:
            df_inv.to_excel(writer, sheet_name="Inventario", index=False)
            df_ventas.to_excel(writer, sheet_name="Ventas", index=False)
            df_detalles.to_excel(
                writer, sheet_name="Detalle Ventas", index=False
            )

            for ws in writer.book.worksheets:
                ws.freeze_panes = "A2"
                ws.sheet_view.showGridLines = False

                for c in ws[1]:
                    c.fill = PatternFill("solid", fgColor="16324F")
                    c.font = Font(bold=True, color="FFFFFF")
                    c.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )

                for col in ws.columns:
                    letra = get_column_letter(col[0].column)
                    ancho = max(
                        [
                            len(str(c.value))
                            if c.value is not None else 0
                            for c in col
                        ] + [12]
                    )
                    ws.column_dimensions[letra].width = min(
                        max(ancho + 3, 14), 35
                    )

        os.replace(temporal, DB)

    except PermissionError:
        if temporal.exists():
            temporal.unlink(missing_ok=True)
        raise PermissionError(
            "Cierra frappes_database.xlsx en Excel y vuelve a guardar."
        )


def obtener_detalles():
    filas = []

    for v in st.session_state.ventas:
        if v.get("estado") != "Completada":
            continue

        for item in v.get("detalles", []):
            filas.append({
                "ID_Pedido": v["id"],
                "Fecha": v["fecha"],
                "Hora": v["hora"],
                "Hora_Entera": str(v["hora"])[:2] + ":00",
                "Cliente": v["cliente"],
                "Modalidad": v["servicio"],
                "Metodo_Pago": v["pago"],
                "Producto": item["producto"],
                "Categoria": item["categoria"],
                "Cantidad": item["cantidad"],
                "Precio_Unitario": item["precio"],
                "Costo_Unitario": item["costo"],
                "Ingreso_Total": item["subtotal"],
                "Costo_Total": item["costo_total"],
                "Ganancia_Neta": item["ganancia_item"]
            })

    return pd.DataFrame(filas, columns=[
        "ID_Pedido", "Fecha", "Hora", "Hora_Entera",
        "Cliente", "Modalidad", "Metodo_Pago", "Producto",
        "Categoria", "Cantidad", "Precio_Unitario",
        "Costo_Unitario", "Ingreso_Total", "Costo_Total",
        "Ganancia_Neta"
    ])


def estilo_figura(fig):
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color="#e2e8f0"),
        margin=dict(l=15, r=15, t=50, b=20),
        legend=dict(orientation="h", yanchor="bottom", y=1.02),
        hoverlabel=dict(
            bgcolor="#172554",
            font_color="white"
        )
    )
    fig.update_xaxes(showgrid=False)
    fig.update_yaxes(gridcolor="#334155")
    return fig


# =========================================================
# 4. CARGAR SESIÓN
# =========================================================

if "inicializado" not in st.session_state:
    inv, ventas, siguiente = cargar_db()

    st.session_state.inventario = inv
    st.session_state.ventas = ventas
    st.session_state.siguiente_id = siguiente
    st.session_state.carrito = []
    st.session_state.inicializado = True

fecha_hora = ahora()

st.title("🥤 FRAPPÉS POS")
st.caption(
    "Sistema de punto de venta y analítica empresarial · "
    + fecha_hora.strftime("%d/%m/%Y %H:%M:%S")
    + " · Bolivia GMT-4"
)

st.sidebar.title("🥤 MENÚ PRINCIPAL")

menu = st.sidebar.radio(
    "Selecciona un módulo",
    [
        "🛒 Registrar venta",
        "📦 Inventario",
        "📊 Dashboard ejecutivo",
        "🗄️ Base de datos Excel",
        "📤 Exportar reporte BI"
    ]
)

st.sidebar.divider()
st.sidebar.caption("💾 Almacenamiento local")
st.sidebar.code(str(DB))


# =========================================================
# 5. REGISTRAR VENTA
# =========================================================

if menu == "🛒 Registrar venta":
    st.header("🛒 Punto de venta")

    c1, c2, c3 = st.columns([2, 1, 1])

    cliente = c1.text_input("Nombre del cliente")
    modalidad = c2.selectbox(
        "Modalidad", ["Para Llevar", "Para Comer Aquí"]
    )
    metodo_pago = c3.selectbox(
        "Método de pago",
        ["Efectivo", "QR / Transferencia", "Tarjeta"]
    )

    notas = st.text_input("Notas de preparación")

    categoria = st.radio(
        "Categoría de productos",
        ["Todas", "Frappés", "Extras"],
        horizontal=True
    )

    productos = [
        nombre for nombre, p in st.session_state.inventario.items()
        if categoria == "Todas" or p["categoria"] == categoria
    ]

    if productos:
        c1, c2, c3 = st.columns([3, 1, 1])

        producto = c1.selectbox("Producto", productos)
        info = st.session_state.inventario[producto]

        c2.metric("Precio", f"Bs. {info['precio']:.2f}")
        c2.caption(f"Stock: {info['stock']}")

        cantidad = c3.number_input(
            "Cantidad",
            min_value=1,
            max_value=max(1, int(info["stock"])),
            value=1,
            step=1,
            disabled=info["stock"] <= 0
        )

        if st.button(
            "➕ Añadir al carrito", use_container_width=True
        ):
            en_carrito = sum(
                x["cantidad"] for x in st.session_state.carrito
                if x["producto"] == producto
            )

            if info["stock"] <= 0:
                st.error("Este producto está agotado.")
            elif en_carrito + cantidad > info["stock"]:
                st.error("La cantidad supera el stock disponible.")
            else:
                st.session_state.carrito.append({
                    "producto": producto,
                    "categoria": info["categoria"],
                    "cantidad": int(cantidad),
                    "precio": info["precio"],
                    "costo": info["costo"],
                    "subtotal": info["precio"] * cantidad,
                    "costo_total": info["costo"] * cantidad,
                    "ganancia_item": (
                        info["precio"] - info["costo"]
                    ) * cantidad
                })
                st.success("Producto añadido.")
                st.rerun()

    if st.session_state.carrito:
        st.divider()
        st.subheader("🛒 Carrito de compra")

        df_carrito = pd.DataFrame(st.session_state.carrito)

        st.dataframe(
            df_carrito[
                ["producto", "cantidad", "precio", "subtotal"]
            ],
            use_container_width=True,
            hide_index=True
        )

        total = sum(x["subtotal"] for x in st.session_state.carrito)
        costo_total = sum(
            x["costo_total"] for x in st.session_state.carrito
        )

        a, b, c = st.columns(3)
        a.metric("Total a cobrar", f"Bs. {total:.2f}")
        b.metric("Costo estimado", f"Bs. {costo_total:.2f}")
        c.metric("Utilidad estimada", f"Bs. {total-costo_total:.2f}")

        a, b = st.columns(2)

        if a.button(
            "✅ Confirmar venta", use_container_width=True
        ):
            if not cliente.strip():
                st.warning("Ingresa el nombre del cliente.")
            else:
                insuficiente = any(
                    st.session_state.inventario[x["producto"]]["stock"]
                    < x["cantidad"]
                    for x in st.session_state.carrito
                )

                if insuficiente:
                    st.error("Stock insuficiente para confirmar.")
                else:
                    venta = {
                        "id": st.session_state.siguiente_id,
                        "fecha": fecha_hora.strftime("%Y-%m-%d"),
                        "hora": fecha_hora.strftime("%H:%M:%S"),
                        "cliente": cliente.strip(),
                        "servicio": modalidad,
                        "pago": metodo_pago,
                        "notas": notas or "Sin especificación",
                        "detalles": [
                            item.copy()
                            for item in st.session_state.carrito
                        ],
                        "total_venta": total,
                        "costo_total": costo_total,
                        "ganancia_neta": total - costo_total,
                        "estado": "Completada"
                    }

                    for item in st.session_state.carrito:
                        st.session_state.inventario[
                            item["producto"]
                        ]["stock"] -= item["cantidad"]

                    st.session_state.ventas.append(venta)
                    st.session_state.siguiente_id += 1
                    st.session_state.carrito = []

                    try:
                        guardar_db()
                        st.success("Venta registrada y guardada en Excel.")
                        st.rerun()
                    except Exception as e:
                        st.error(
                            f"No se pudo guardar en Excel: {e}. "
                            "Revisa el archivo de base de datos."
                        )

        if b.button("🗑️ Vaciar carrito", use_container_width=True):
            st.session_state.carrito = []
            st.rerun()

    st.divider()
    st.subheader("📋 Historial de ventas")

    if st.session_state.ventas:
        df_historial = pd.DataFrame([
            {
                "ID": v["id"],
                "Fecha": v["fecha"],
                "Hora": v["hora"],
                "Cliente": v["cliente"],
                "Modalidad": v["servicio"],
                "Pago": v["pago"],
                "Total (Bs.)": v["total_venta"],
                "Estado": v["estado"]
            }
            for v in st.session_state.ventas
        ])

        st.dataframe(
            df_historial.sort_values("ID", ascending=False),
            use_container_width=True,
            hide_index=True
        )

        ids = [v["id"] for v in st.session_state.ventas]

        with st.expander("🗑️ Eliminar una venta"):
            id_eliminar = st.selectbox(
                "Selecciona el pedido",
                ids,
                key="id_eliminar"
            )

            confirmar = st.checkbox(
                "Confirmo eliminar la venta y restituir el stock."
            )

            if st.button("Eliminar venta seleccionada"):
                if not confirmar:
                    st.warning("Confirma la operación primero.")
                else:
                    venta = next(
                        v for v in st.session_state.ventas
                        if v["id"] == id_eliminar
                    )

                    if venta["estado"] == "Completada":
                        for item in venta["detalles"]:
                            nombre = item["producto"]
                            if nombre in st.session_state.inventario:
                                st.session_state.inventario[nombre][
                                    "stock"
                                ] += item["cantidad"]

                    st.session_state.ventas = [
                        v for v in st.session_state.ventas
                        if v["id"] != id_eliminar
                    ]

                    try:
                        guardar_db()
                        st.success("Venta eliminada y stock restituido.")
                        st.rerun()
                    except Exception as e:
                        st.error(f"Error al guardar: {e}")
    else:
        st.info("Todavía no hay ventas registradas.")


# =========================================================
# 6. INVENTARIO
# =========================================================

elif menu == "📦 Inventario":
    st.header("📦 Control y gestión de inventario")

    df_inv = pd.DataFrame([
        {
            "Producto": n,
            "Categoría": p["categoria"],
            "Precio (Bs.)": p["precio"],
            "Costo (Bs.)": p["costo"],
            "Stock": p["stock"],
            "Utilidad por unidad": p["precio"] - p["costo"]
        }
        for n, p in st.session_state.inventario.items()
    ])

    agotados = int((df_inv["Stock"] == 0).sum())
    bajo_stock = int(
        ((df_inv["Stock"] > 0) & (df_inv["Stock"] <= 5)).sum()
    )

    a, b, c = st.columns(3)
    a.metric("Productos", len(df_inv))
    b.metric("Con stock bajo", bajo_stock)
    c.metric("Agotados", agotados)

    st.dataframe(
        df_inv.style.format({
            "Precio (Bs.)": "Bs. {:.2f}",
            "Costo (Bs.)": "Bs. {:.2f}",
            "Utilidad por unidad": "Bs. {:.2f}"
        }),
        use_container_width=True,
        hide_index=True
    )

    tab1, tab2 = st.tabs(["✏️ Editar producto", "➕ Nuevo producto"])

    with tab1:
        nombre = st.selectbox(
            "Producto a editar",
            list(st.session_state.inventario.keys())
        )

        p = st.session_state.inventario[nombre]

        a, b, c = st.columns(3)

        nuevo_stock = a.number_input(
            "Stock disponible",
            min_value=0,
            value=int(p["stock"]),
            step=1
        )

        nuevo_precio = b.number_input(
            "Precio de venta (Bs.)",
            min_value=0.0,
            value=float(p["precio"]),
            step=0.5,
            key="editar_precio"
        )

        nuevo_costo = c.number_input(
            "Costo (Bs.)",
            min_value=0.0,
            value=float(p["costo"]),
            step=0.5,
            key="editar_costo"
        )

        nueva_categoria = st.selectbox(
            "Categoría del producto",
            ["Frappés", "Extras"],
            index=0 if p["categoria"] == "Frappés" else 1,
            key="editar_categoria"
        )

        if st.button("💾 Guardar modificaciones"):
            p.update({
                "stock": int(nuevo_stock),
                "precio": float(nuevo_precio),
                "costo": float(nuevo_costo),
                "categoria": nueva_categoria
            })

            try:
                guardar_db()
                st.success("Producto actualizado.")
                st.rerun()
            except Exception as e:
                st.error(f"Error al guardar: {e}")

    with tab2:
        nombre_nuevo = st.text_input(
            "Nombre del nuevo producto"
        )

        cat_nueva = st.selectbox(
            "Categoría", ["Frappés", "Extras"], key="nueva_categoria"
        )

        precio_nuevo = st.number_input(
            "Precio de venta",
            min_value=0.0,
            value=20.0,
            key="nuevo_precio"
        )

        costo_nuevo = st.number_input(
            "Costo de producción",
            min_value=0.0,
            value=7.0,
            key="nuevo_costo"
        )

        stock_nuevo = st.number_input(
            "Stock inicial",
            min_value=0,
            value=20,
            step=1,
            key="nuevo_stock"
        )

        if st.button("➕ Crear producto"):
            nombre_nuevo = nombre_nuevo.strip()

            if not nombre_nuevo:
                st.error("Escribe un nombre.")
            elif nombre_nuevo in st.session_state.inventario:
                st.error("El producto ya existe.")
            else:
                st.session_state.inventario[nombre_nuevo] = {
                    "categoria": cat_nueva,
                    "precio": float(precio_nuevo),
                    "costo": float(costo_nuevo),
                    "stock": int(stock_nuevo)
                }

                try:
                    guardar_db()
                    st.success("Producto creado.")
                    st.rerun()
                except Exception as e:
                    st.error(f"Error al guardar: {e}")


# =========================================================
# 7. DASHBOARD EJECUTIVO
# =========================================================

elif menu == "📊 Dashboard ejecutivo":
    st.header("📊 Dashboard ejecutivo de ventas")

    df = obtener_detalles()
    ventas_ok = [
        v for v in st.session_state.ventas
        if v["estado"] == "Completada"
    ]

    if df.empty:
        st.info(
            "Registra una venta para visualizar los gráficos e indicadores."
        )
    else:
        ingresos = sum(v["total_venta"] for v in ventas_ok)
        costos = sum(v["costo_total"] for v in ventas_ok)
        utilidad = sum(v["ganancia_neta"] for v in ventas_ok)
        unidades = int(df["Cantidad"].sum())
        margen = utilidad / ingresos if ingresos else 0

        # FILTRO DE FECHAS
        fechas = sorted(df["Fecha"].dropna().astype(str).unique())

        with st.expander("🔎 Filtrar el dashboard por fecha"):
            fechas_seleccionadas = st.multiselect(
                "Fechas incluidas",
                fechas,
                default=fechas
            )

        if fechas_seleccionadas:
            df = df[df["Fecha"].isin(fechas_seleccionadas)]
            ids_filtrados = set(df["ID_Pedido"])
            ventas_filtradas = [
                v for v in ventas_ok if v["id"] in ids_filtrados
            ]
        else:
            df = df.iloc[0:0]
            ventas_filtradas = []

        ingresos = sum(v["total_venta"] for v in ventas_filtradas)
        costos = sum(v["costo_total"] for v in ventas_filtradas)
        utilidad = sum(v["ganancia_neta"] for v in ventas_filtradas)
        unidades = int(df["Cantidad"].sum()) if not df.empty else 0
        margen = utilidad / ingresos if ingresos else 0

        k1, k2, k3, k4, k5, k6 = st.columns(6)
        k1.metric("🧾 Pedidos", len(ventas_filtradas))
        k2.metric("💰 Ingresos", f"Bs. {ingresos:,.2f}")
        k3.metric("📉 Costos", f"Bs. {costos:,.2f}")
        k4.metric("💚 Utilidad", f"Bs. {utilidad:,.2f}")
        k5.metric("📈 Margen", f"{margen:.1%}")
        k6.metric("🥤 Unidades", unidades)

        st.divider()

        col1, col2 = st.columns(2)

        with col1:
            st.subheader("1. Ingresos y utilidad por pedido")

            q = df.groupby("ID_Pedido", as_index=False).agg(
                Ingresos=("Ingreso_Total", "sum"),
                Utilidad=("Ganancia_Neta", "sum")
            )

            fig = go.Figure()

            fig.add_bar(
                x=q["ID_Pedido"],
                y=q["Ingresos"],
                name="Ingresos",
                marker_color="#22D3EE",
                text=q["Ingresos"].map(lambda x: f"Bs. {x:.2f}"),
                textposition="outside"
            )

            fig.add_scatter(
                x=q["ID_Pedido"],
                y=q["Utilidad"],
                name="Utilidad",
                mode="lines+markers+text",
                text=q["Utilidad"].map(lambda x: f"Bs. {x:.2f}"),
                textposition="top center",
                line=dict(color="#4ADE80", width=3)
            )

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        with col2:
            st.subheader("2. Recaudación por método de pago")

            q = df.groupby("Metodo_Pago", as_index=False).agg(
                Ingresos=("Ingreso_Total", "sum"),
                Pedidos=("ID_Pedido", "nunique")
            )

            fig = px.pie(
                q,
                names="Metodo_Pago",
                values="Ingresos",
                hole=0.48,
                color_discrete_sequence=[
                    "#A78BFA", "#22D3EE", "#FB7185"
                ]
            )

            fig.update_traces(
                textinfo="label+percent",
                customdata=q[["Pedidos"]],
                hovertemplate=(
                    "%{label}<br>Ingresos: Bs. %{value:.2f}"
                    "<br>Pedidos: %{customdata[0]}<extra></extra>"
                )
            )

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        st.divider()

        col3, col4 = st.columns(2)

        with col3:
            st.subheader("3. Ingresos frente a costos")

            q = df.groupby("Producto", as_index=False).agg(
                Ingresos=("Ingreso_Total", "sum"),
                Costos=("Costo_Total", "sum")
            )

            fig = go.Figure()

            fig.add_bar(
                x=q["Producto"], y=q["Ingresos"],
                name="Ingresos", marker_color="#38BDF8",
                text=q["Ingresos"].map(lambda x: f"Bs. {x:.0f}"),
                textposition="outside"
            )

            fig.add_bar(
                x=q["Producto"], y=q["Costos"],
                name="Costos", marker_color="#FB7185",
                text=q["Costos"].map(lambda x: f"Bs. {x:.0f}"),
                textposition="outside"
            )

            fig.update_layout(barmode="group")

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        with col4:
            st.subheader("4. Unidades vendidas por producto")

            q = df.groupby("Producto", as_index=False).agg(
                Unidades=("Cantidad", "sum")
            ).sort_values("Unidades")

            fig = px.bar(
                q,
                x="Unidades",
                y="Producto",
                orientation="h",
                text="Unidades",
                color="Unidades",
                color_continuous_scale="Teal"
            )

            fig.update_traces(textposition="outside")

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        st.divider()

        col5, col6 = st.columns(2)

        with col5:
            st.subheader("5. Ingresos por hora")

            q = df.groupby("Hora_Entera", as_index=False).agg(
                Ingresos=("Ingreso_Total", "sum"),
                Pedidos=("ID_Pedido", "nunique")
            ).sort_values("Hora_Entera")

            fig = px.bar(
                q,
                x="Hora_Entera",
                y="Ingresos",
                text="Ingresos",
                color="Ingresos",
                color_continuous_scale="Plasma",
                hover_data=["Pedidos"]
            )

            fig.update_traces(
                texttemplate="Bs. %{text:.2f}",
                textposition="outside"
            )

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        with col6:
            st.subheader("6. Participación por categoría")

            q = df.groupby("Categoria", as_index=False).agg(
                Ingresos=("Ingreso_Total", "sum"),
                Unidades=("Cantidad", "sum")
            )

            fig = px.pie(
                q,
                names="Categoria",
                values="Ingresos",
                hole=0.4,
                color_discrete_sequence=[
                    "#F59E0B", "#34D399", "#818CF8"
                ]
            )

            fig.update_traces(
                textinfo="label+percent",
                hovertemplate=(
                    "%{label}<br>Ingresos: Bs. %{value:.2f}"
                    "<extra></extra>"
                )
            )

            st.plotly_chart(
                estilo_figura(fig), use_container_width=True
            )

        st.divider()
        st.subheader("📋 Tabla de rendimiento por producto")

        tabla = df.groupby(
            ["Producto", "Categoria"], as_index=False
        ).agg(
            Unidades=("Cantidad", "sum"),
            Ingresos=("Ingreso_Total", "sum"),
            Costos=("Costo_Total", "sum"),
            Utilidad=("Ganancia_Neta", "sum")
        )

        tabla["Margen (%)"] = (
            tabla["Utilidad"]
            / tabla["Ingresos"].replace(0, 1) * 100
        )

        st.dataframe(
            tabla.style.format({
                "Ingresos": "Bs. {:.2f}",
                "Costos": "Bs. {:.2f}",
                "Utilidad": "Bs. {:.2f}",
                "Margen (%)": "{:.2f}%"
            }),
            use_container_width=True,
            hide_index=True
        )

        st.subheader("📅 Resumen diario")

        diario = df.groupby("Fecha", as_index=False).agg(
            Pedidos=("ID_Pedido", "nunique"),
            Unidades=("Cantidad", "sum"),
            Ingresos=("Ingreso_Total", "sum"),
            Costos=("Costo_Total", "sum"),
            Utilidad=("Ganancia_Neta", "sum")
        )

        st.dataframe(
            diario,
            use_container_width=True,
            hide_index=True
        )


# =========================================================
# 8. BASE DE DATOS EXCEL
# =========================================================

elif menu == "🗄️ Base de datos Excel":
    st.header("🗄️ Base de datos local")

    st.write(
        "El sistema guarda el inventario, las ventas y el detalle "
        "de cada pedido en un libro Excel local."
    )

    st.code(str(DB))

    tab1, tab2, tab3 = st.tabs([
        "📦 Inventario", "💰 Ventas", "🧾 Detalle de ventas"
    ])

    with tab1:
        st.dataframe(
            pd.DataFrame([
                {"Producto": n, **p}
                for n, p in st.session_state.inventario.items()
            ]),
            use_container_width=True,
            hide_index=True
        )

    with tab2:
        st.dataframe(
            pd.DataFrame([
                {
                    "ID": v["id"],
                    "Fecha": v["fecha"],
                    "Hora": v["hora"],
                    "Cliente": v["cliente"],
                    "Total": v["total_venta"],
                    "Costo": v["costo_total"],
                    "Utilidad": v["ganancia_neta"],
                    "Estado": v["estado"]
                }
                for v in st.session_state.ventas
            ]),
            use_container_width=True,
            hide_index=True
        )

    with tab3:
        st.dataframe(
            obtener_detalles(),
            use_container_width=True,
            hide_index=True
        )

    if DB.exists():
        st.download_button(
            "⬇️ Descargar copia de la base de datos",
            data=DB.read_bytes(),
            file_name="frappes_database_copia.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True
        )

    if st.button("💾 Guardar base de datos ahora"):
        try:
            guardar_db()
            st.success("Base de datos guardada correctamente.")
        except Exception as e:
            st.error(f"No se pudo guardar: {e}")


# =========================================================
# 9. EXPORTAR REPORTE BI PROFESIONAL
# =========================================================

elif menu == "📤 Exportar reporte BI":
    st.header("📤 Suite de inteligencia de negocios")

    df = obtener_detalles()

    st.write(
        "Crea un Excel con dashboard, gráficos, indicadores, "
        "inventario, ventas y hojas analíticas."
    )

    if st.button(
        "🚀 Generar Excel profesional",
        use_container_width=True
    ):
        salida = BytesIO()

        if not df.empty:
            productos = df.groupby(
                ["Producto", "Categoria"], as_index=False
            ).agg(
                Unidades=("Cantidad", "sum"),
                Ingresos=("Ingreso_Total", "sum"),
                Costos=("Costo_Total", "sum"),
                Utilidad=("Ganancia_Neta", "sum")
            )

            productos["Margen_%"] = (
                productos["Utilidad"]
                / productos["Ingresos"].replace(0, 1)
            )

            total_ing = productos["Ingresos"].sum()

            productos["Participacion_%"] = (
                productos["Ingresos"] / total_ing
                if total_ing else 0
            )

            pagos = df.groupby(
                "Metodo_Pago", as_index=False
            ).agg(
                Pedidos=("ID_Pedido", "nunique"),
                Ingresos=("Ingreso_Total", "sum")
            )

            horas = df.groupby(
                "Hora_Entera", as_index=False
            ).agg(
                Pedidos=("ID_Pedido", "nunique"),
                Unidades=("Cantidad", "sum"),
                Ingresos=("Ingreso_Total", "sum")
            ).sort_values("Hora_Entera")

            categorias = df.groupby(
                "Categoria", as_index=False
            ).agg(
                Unidades=("Cantidad", "sum"),
                Ingresos=("Ingreso_Total", "sum")
            )

            diario = df.groupby(
                "Fecha", as_index=False
            ).agg(
                Pedidos=("ID_Pedido", "nunique"),
                Unidades=("Cantidad", "sum"),
                Ingresos=("Ingreso_Total", "sum"),
                Costos=("Costo_Total", "sum"),
                Utilidad=("Ganancia_Neta", "sum")
            )

        else:
            productos = pd.DataFrame(columns=[
                "Producto", "Categoria", "Unidades", "Ingresos",
                "Costos", "Utilidad", "Margen_%", "Participacion_%"
            ])
            pagos = pd.DataFrame(
                columns=["Metodo_Pago", "Pedidos", "Ingresos"]
            )
            horas = pd.DataFrame(
                columns=["Hora_Entera", "Pedidos", "Unidades", "Ingresos"]
            )
            categorias = pd.DataFrame(
                columns=["Categoria", "Unidades", "Ingresos"]
            )
            diario = pd.DataFrame(
                columns=[
                    "Fecha", "Pedidos", "Unidades", "Ingresos",
                    "Costos", "Utilidad"
                ]
            )

        inventario_df = pd.DataFrame([
            {
                "Producto": n,
                "Categoria": p["categoria"],
                "Precio": p["precio"],
                "Costo": p["costo"],
                "Stock": p["stock"],
                "Utilidad_Unit": p["precio"] - p["costo"]
            }
            for n, p in st.session_state.inventario.items()
        ])

        ventas_df = pd.DataFrame([
            {
                "ID": v["id"],
                "Fecha": v["fecha"],
                "Hora": v["hora"],
                "Cliente": v["cliente"],
                "Modalidad": v["servicio"],
                "Metodo_Pago": v["pago"],
                "Notas": v.get("notas", ""),
                "Total_Venta": v["total_venta"],
                "Costo_Total": v["costo_total"],
                "Ganancia_Neta": v["ganancia_neta"],
                "Estado": v["estado"]
            }
            for v in st.session_state.ventas
        ])

        total_ventas = (
            ventas_df["Total_Venta"].sum()
            if not ventas_df.empty else 0
        )
        total_costos = (
            ventas_df["Costo_Total"].sum()
            if not ventas_df.empty else 0
        )
        total_utilidad = (
            ventas_df["Ganancia_Neta"].sum()
            if not ventas_df.empty else 0
        )
        total_unidades = int(df["Cantidad"].sum()) if not df.empty else 0

        with pd.ExcelWriter(salida, engine="openpyxl") as writer:
            productos.to_excel(writer, sheet_name="Productos", index=False)
            pagos.to_excel(writer, sheet_name="Metodos de Pago", index=False)
            horas.to_excel(writer, sheet_name="Horas Pico", index=False)
            categorias.to_excel(writer, sheet_name="Categorias", index=False)
            diario.to_excel(writer, sheet_name="Resumen Diario", index=False)
            inventario_df.to_excel(writer, sheet_name="Inventario", index=False)
            ventas_df.to_excel(writer, sheet_name="Ventas", index=False)
            df.to_excel(writer, sheet_name="Detalle Ventas", index=False)

            wb = writer.book

            # FORMATO DE LAS HOJAS
            for ws in wb.worksheets:
                ws.freeze_panes = "A2"
                ws.sheet_view.showGridLines = False

                for cell in ws[1]:
                    cell.fill = PatternFill(
                        "solid", fgColor="16324F"
                    )
                    cell.font = Font(
                        bold=True, color="FFFFFF"
                    )
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )

                for col in ws.columns:
                    letra = get_column_letter(col[0].column)
                    ancho = max(
                        [
                            len(str(c.value))
                            if c.value is not None else 0
                            for c in col
                        ] + [12]
                    )
                    ws.column_dimensions[letra].width = min(
                        max(ancho + 3, 14), 35
                    )

                if ws.max_row >= 2 and ws.max_column >= 1:
                    nombre_tabla = "Tabla" + "".join(
                        ch for ch in ws.title if ch.isalnum()
                    )[:20]

                    tabla = Table(
                        displayName=nombre_tabla,
                        ref=ws.dimensions
                    )
                    tabla.tableStyleInfo = TableStyleInfo(
                        name="TableStyleMedium2",
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=True,
                        showColumnStripes=False
                    )
                    ws.add_table(tabla)

            # DASHBOARD MAESTRO
            dash = wb.create_sheet("Dashboard Ejecutivo", 0)
            dash.sheet_view.showGridLines = False

            dash.merge_cells("B2:Q2")
            dash["B2"] = "🥤 FRAPPÉS | DASHBOARD EJECUTIVO"
            dash["B2"].fill = PatternFill(
                "solid", fgColor="102A43"
            )
            dash["B2"].font = Font(
                size=18, bold=True, color="FFFFFF"
            )
            dash["B2"].alignment = Alignment(
                horizontal="center", vertical="center"
            )
            dash.row_dimensions[2].height = 42

            tarjetas = [
                ("INGRESOS (Bs.)", total_ventas, "B", "D",
                 "0891B2", "#,##0.00"),
                ("COSTOS (Bs.)", total_costos, "E", "G",
                 "D97706", "#,##0.00"),
                ("UTILIDAD (Bs.)", total_utilidad, "H", "J",
                 "15803D", "#,##0.00"),
                ("MARGEN", total_utilidad / total_ventas
                 if total_ventas else 0, "K", "M",
                 "6D28D9", "0.00%"),
                ("PEDIDOS", len(ventas_df), "N", "O",
                 "BE185D", "#,##0"),
                ("UNIDADES", total_unidades, "P", "Q",
                 "0F766E", "#,##0")
            ]

            for etiqueta, valor, c1, c2, color, formato in tarjetas:
                dash.merge_cells(f"{c1}4:{c2}4")
                dash.merge_cells(f"{c1}5:{c2}6")

                cab = dash[f"{c1}4"]
                cab.value = etiqueta
                cab.fill = PatternFill("solid", fgColor=color)
                cab.font = Font(bold=True, color="FFFFFF")
                cab.alignment = Alignment(horizontal="center")

                celda = dash[f"{c1}5"]
                celda.value = valor
                celda.number_format = formato
                celda.fill = PatternFill("solid", fgColor="EAF2F8")
                celda.font = Font(
                    size=14, bold=True, color="102A43"
                )
                celda.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            # TABLA RESUMEN DE PRODUCTOS EN EL DASHBOARD
            dash["B8"] = "RESUMEN DE RENDIMIENTO POR PRODUCTO"
            dash["B8"].font = Font(bold=True, size=12, color="102A43")

            resumen_cols = [
                "Producto", "Categoria", "Unidades", "Ingresos",
                "Costos", "Utilidad", "Margen_%", "Participacion_%"
            ]

            for j, nombre_col in enumerate(resumen_cols, 2):
                celda = dash.cell(row=9, column=j, value=nombre_col)
                celda.fill = PatternFill("solid", fgColor="16324F")
                celda.font = Font(bold=True, color="FFFFFF")
                celda.alignment = Alignment(horizontal="center")

            for i, fila in enumerate(
                productos[resumen_cols].itertuples(index=False, name=None),
                10
            ):
                for j, valor in enumerate(fila, 2):
                    celda = dash.cell(row=i, column=j, value=valor)
                    celda.border = Border(
                        bottom=Side(style="thin", color="CBD5E1")
                    )

                    if j in (5, 6, 7):
                        celda.number_format = '#,##0.00'
                    elif j in (8, 9):
                        celda.number_format = '0.00%'

            # FUNCIÓN PARA CREAR GRÁFICOS
            def agregar_grafico(
                tipo, titulo, hoja, col_datos,
                col_categoria, ancla, ancho=15, alto=8
            ):
                origen = wb[hoja]

                if origen.max_row < 2:
                    return

                grafico = tipo()
                grafico.title = titulo

                datos = Reference(
                    origen,
                    min_col=col_datos,
                    min_row=1,
                    max_row=origen.max_row
                )

                categorias_ref = Reference(
                    origen,
                    min_col=col_categoria,
                    min_row=2,
                    max_row=origen.max_row
                )

                grafico.add_data(datos, titles_from_data=True)
                grafico.set_categories(categorias_ref)
                grafico.width = ancho
                grafico.height = alto
                grafico.style = 10

                if tipo == BarChart:
                    grafico.type = "col"
                    grafico.legend.position = "b"

                grafico.dataLabels = DataLabelList()
                grafico.dataLabels.showVal = True

                dash.add_chart(grafico, ancla)

            # 6 GRÁFICOS EN EL DASHBOARD DE EXCEL
            agregar_grafico(
                BarChart, "Ingresos por producto (Bs.)",
                "Productos", 4, 1, "B25"
            )

            agregar_grafico(
                BarChart, "Utilidad por producto (Bs.)",
                "Productos", 6, 1, "J25"
            )

            agregar_grafico(
                DoughnutChart, "Recaudación por método de pago",
                "Metodos de Pago", 3, 1, "B42"
            )

            agregar_grafico(
                LineChart, "Ingresos por hora",
                "Horas Pico", 4, 1, "H42"
            )

            agregar_grafico(
                PieChart, "Ingresos por categoría",
                "Categorias", 3, 1, "N42"
            )

            agregar_grafico(
                BarChart, "Unidades por producto",
                "Productos", 3, 1, "B59"
            )

            for col in range(1, 18):
                dash.column_dimensions[
                    get_column_letter(col)
                ].width = 15

        st.session_state.archivo_excel = salida.getvalue()
        st.success("¡Excel profesional generado correctamente!")

    if "archivo_excel" in st.session_state:
        st.download_button(
            "⬇️ DESCARGAR EXCEL CON DASHBOARD Y GRÁFICOS",
            data=st.session_state.archivo_excel,
            file_name=(
                f"Frappes_Dashboard_{ahora():%Y%m%d_%H%M}.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True
        )
